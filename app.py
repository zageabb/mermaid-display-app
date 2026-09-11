from flask import Flask, render_template, abort, request, redirect, url_for, flash, send_file
from diagram_pptx import render_mermaid
from pptx import Presentation
from pptx.util import Inches
import io
import os
import re

app = Flask(__name__)
app.secret_key = 'super_secret_dev_key'

DIAGRAM_DIR = "diagrams"
SUB_DIAGRAM_DIR = os.path.join(DIAGRAM_DIR, "sub_diagrams")
POWERPOINT_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _diagram_path(rel_path):
    """Resolve a diagram path while preventing traversal outside the diagram repository."""
    root = os.path.abspath(DIAGRAM_DIR)
    candidate = os.path.abspath(os.path.join(root, rel_path))
    try:
        inside_root = os.path.commonpath([root, candidate]) == root
    except ValueError:
        inside_root = False
    if not inside_root:
        abort(400)
    return candidate


def assemble_diagram(filepath, is_sub_include=False):
    """Recursively parses and patches %% INCLUDE statements in memory."""
    if not os.path.exists(filepath):
        return None

    with open(filepath, 'r') as f:
        lines = f.readlines()

    assembled_lines = []
    include_pattern = re.compile(r"^\s*%%\s*INCLUDE\s+(.+)$")

    for line in lines:
        # If this is an included sub-file, skip any redundant graph definitions
        if is_sub_include and ("erDiagram" in line or "flowchart" in line):
            continue

        match = include_pattern.match(line)
        if match:
            sub_filename = match.group(1).strip()
            sub_filepath = _diagram_path(sub_filename)

            # Pass True for is_sub_include flag on recursive calls
            sub_content = assemble_diagram(sub_filepath, is_sub_include=True)
            if sub_content:
                assembled_lines.append(f"\n{sub_content}\n")
            else:
                assembled_lines.append(f"%% ERROR: Missing include '{sub_filename}'\n")
        else:
            assembled_lines.append(line)

    return "".join(assembled_lines)


def build_powerpoint(content):
    """Create a 16:9 PPTX containing native editable shapes from Mermaid source."""
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    render_mermaid(content, slide=slide, position="full")
    output = io.BytesIO()
    deck.save(output)
    output.seek(0)
    return output


@app.route("/")
def index():
    os.makedirs(SUB_DIAGRAM_DIR, exist_ok=True)

    # 1. Gather all physical files from disk
    master_files = [f for f in os.listdir(DIAGRAM_DIR) if f.endswith('.mmd')]
    sub_files = [f for f in os.listdir(SUB_DIAGRAM_DIR) if f.endswith('.mmd')]

    # Track which sub-diagrams have been assigned to a master container
    assigned_subs = set()
    ordered_dashboard_list = []

    # Regex to extract the file name from: %% INCLUDE sub_diagrams/filename.mmd
    include_pattern = re.compile(r"^\s*%%\s*INCLUDE\s+sub_diagrams/(.+)$")

    # 2. Map Sub-Documents to their parent Master Documents
    for master in master_files:
        # Add the Master record to our dashboard array
        ordered_dashboard_list.append({
            "name": master,
            "type": "Master",
            "rel_path": master,
            "is_child": False
        })

        master_path = os.path.join(DIAGRAM_DIR, master)
        try:
            with open(master_path, 'r') as f:
                for line in f:
                    match = include_pattern.match(line)
                    if match:
                        sub_name = match.group(1).strip()
                        if sub_name in sub_files:
                            ordered_dashboard_list.append({
                                "name": sub_name,
                                "type": "Sub-Component",
                                "rel_path": f"sub_diagrams/{sub_name}",
                                "is_child": True  # Flag to trigger visual indentation
                            })
                            assigned_subs.add(sub_name)
        except Exception as e:
            print(f"Error parsing master mapping for {master}: {e}")

    # 3. Clean up orphaned sub-diagrams (files not explicitly listed in any master)
    orphans = [s for s in sub_files if s not in assigned_subs]
    if orphans:
        for orphan in orphans:
            ordered_dashboard_list.append({
                "name": orphan,
                "type": "Sub-Component (Unlinked)",
                "rel_path": f"sub_diagrams/{orphan}",
                "is_child": False
            })

    return render_template("index.html", files=ordered_dashboard_list)


@app.route("/view/<path:rel_path>")
def view(rel_path):
    file_path = _diagram_path(rel_path)
    content = assemble_diagram(file_path, is_sub_include=False)

    if content is None:
        abort(404)
    return render_template("view.html", filename=rel_path, content=content)


@app.route("/export-pptx/<path:rel_path>")
def export_pptx(rel_path):
    """Download the assembled Mermaid diagram as an editable PowerPoint slide."""
    file_path = _diagram_path(rel_path)
    content = assemble_diagram(file_path, is_sub_include=False)
    if content is None:
        abort(404)
    try:
        output = build_powerpoint(content)
    except Exception as exc:
        abort(422, description=f"PowerPoint export failed: {exc}")

    download_name = f"{os.path.splitext(os.path.basename(rel_path))[0]}.pptx"
    return send_file(
        output,
        mimetype=POWERPOINT_MIME,
        as_attachment=True,
        download_name=download_name,
    )


@app.route("/edit/<path:rel_path>", methods=['GET', 'POST'])
def edit(rel_path):
    file_path = _diagram_path(rel_path)

    if request.method == 'POST':
        new_code = request.form.get('code')
        try:
            with open(file_path, 'w') as f:
                f.write(new_code)
            flash(f"Successfully saved {rel_path}!")
            return redirect(url_for('index'))
        except Exception as e:
            flash(f"Error saving file: {str(e)}")

    if not os.path.exists(file_path):
        abort(404)

    with open(file_path, 'r') as f:
        raw_content = f.read()

    return render_template("edit.html", filename=rel_path, content=raw_content)


if __name__ == "__main__":
    app.run(debug=True, host='0.0.0.0', port='5012')
